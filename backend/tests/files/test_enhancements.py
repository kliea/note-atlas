"""Files F6: lab enhancement activities, each on and off."""

from __future__ import annotations

import io
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import docx
import pytest
from fastapi.testclient import TestClient

from app.files import answer, repo, service
from app.files.extract import extract_text

# --- 1. Word (.docx) ----------------------------------------------------------------


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Methodology", level=1)
    document.add_paragraph("We surveyed 120 respondents.")
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Group"
    table.rows[0].cells[1].text = "Size"
    merged = table.rows[1].cells[0].merge(table.rows[1].cells[1])
    merged.text = "Students"
    table.rows[1].cells[2].text = "120"
    document.add_paragraph("Results follow the table.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_is_accepted_and_read_in_document_order(tmp_path: Path) -> None:
    data = _docx_bytes()
    assert repo.classify_upload("study.docx", data) == repo.MIME_DOCX
    path = tmp_path / "study.docx"
    path.write_bytes(data)
    pages = extract_text(path, repo.MIME_DOCX)
    assert [p.page_no for p in pages] == [1]
    assert pages[0].text == (
        "Methodology\nWe surveyed 120 respondents.\nGroup | Size\nStudents | 120\nResults follow the table."
    )


def test_docx_builds_and_is_retrieved_with_its_name() -> None:
    folder = service.create_folder("Word")
    service.add_files(folder.id, [("study.docx", _docx_bytes())])
    assert service.build_folder(folder.id).indexed == 1
    document = service.list_documents(folder.id)[0]
    assert (document.status, document.page_count, document.chunk_count) == ("indexed", 1, 1)


def test_broken_docx_is_rejected_at_upload() -> None:
    with pytest.raises(repo.UnsupportedFile, match="not a readable Word document"):
        repo.classify_upload("broken.docx", b"PK\x03\x04garbage")
    with pytest.raises(repo.UnsupportedFile, match="not a Word document"):
        repo.classify_upload("fake.docx", b"hello")


def test_docx_off_rejects_word_files(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FILES_DOCX", "0")
    with pytest.raises(repo.UnsupportedFile, match=r"only \.pdf, \.txt files"):
        repo.classify_upload("study.docx", _docx_bytes())
    assert repo.accepted_suffixes() == [".pdf", ".txt"]


# --- 2. Text encoding fallback ---------------------------------------------------------

CP1252 = "Café – résumé “quoted”".encode("cp1252")


def test_non_utf8_text_falls_back_to_windows_1252_and_reports_it(client: TestClient) -> None:
    folder = client.post("/files/folders", json={"name": "Notes"}).json()
    body = client.post(
        f"/files/folders/{folder['id']}/documents",
        files=[("files", ("legacy.txt", CP1252, "text/plain")), ("files", ("modern.txt", "Café".encode(), "text/plain"))],
    ).json()
    assert [(d["filename"], d["encoding"]) for d in body["added"]] == [
        ("legacy.txt", "windows-1252"),
        ("modern.txt", "utf-8"),
    ]
    preview = client.get(f"/files/documents/{body['added'][0]['id']}/chunks").json()
    assert preview["chunks"][0]["text"] == "Café – résumé “quoted”"


def test_fallback_off_keeps_utf8_only(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FILES_TXT_FALLBACK", "0")
    with pytest.raises(repo.UnsupportedFile, match="not UTF-8"):
        repo.classify_upload("legacy.txt", CP1252)


def test_utf8_bom_is_dropped_and_old_rows_read_as_utf8(tmp_path: Path) -> None:
    path = tmp_path / "bom.txt"
    path.write_bytes("﻿hello".encode("utf-8"))
    assert extract_text(path, repo.MIME_TEXT)[0].text == "hello"
    assert extract_text(path, "text/plain")[0].text == "hello"  # rows stored before F6
    assert repo.text_encoding("text/plain") == "utf-8"
    assert repo.text_encoding(repo.MIME_PDF) is None


# --- shared fakes for Ask ----------------------------------------------------------------


class FakeChat:
    def __init__(self) -> None:
        self.reply = ["Grounded answer."]
        self.calls: list[dict[str, Any]] = []

    def __call__(self, messages: list[dict[str, str]], *, model: str | None = None, **_: Any) -> Iterator[str]:
        self.calls.append({"prompt": messages[0]["content"], "model": model})
        yield from self.reply


@pytest.fixture
def chat(monkeypatch: pytest.MonkeyPatch) -> FakeChat:
    fake = FakeChat()
    monkeypatch.setattr("app.llm.chat_stream", fake)
    return fake


def _built(name: str, *files: tuple[str, bytes]) -> int:
    folder = service.create_folder(name)
    service.add_files(folder.id, list(files))
    service.build_folder(folder.id)
    return folder.id


TXT = ("Retrieval grounds answers in documents. " * 30).encode()


# --- 3. Document names on retrieved chunks / source line ------------------------------------


def test_answer_carries_a_source_line(chat: FakeChat) -> None:
    folder_id = _built("Lectures", ("week1.txt", TXT))
    service.ask("q", folder_ids=[folder_id])
    assert service.chat_history()[0].sources == ["week1.txt · p.1"]


def test_refusal_has_no_source_line(chat: FakeChat) -> None:
    chat.reply = [answer.REFUSAL]
    _built("Lectures", ("week1.txt", TXT))
    service.ask("off topic")
    assert service.chat_history()[0].sources is None


def test_source_line_off(chat: FakeChat, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FILES_SHOW_SOURCES", "0")
    _built("Lectures", ("week1.txt", TXT))
    service.ask("q")
    assert service.chat_history()[0].sources is None


def test_source_line_reaches_the_http_done_event(chat: FakeChat, client: TestClient) -> None:
    _built("Lectures", ("week1.txt", TXT))
    body = client.post("/files/ask", json={"question": "q"}).text
    done = [b for b in body.split("\n\n") if b.startswith("event: done")][0]
    assert '"sources": ["week1.txt \\u00b7 p.1"]' in done


# --- 4. Hybrid search -----------------------------------------------------------------------

FILLER = ("filler text " * 1100).encode()  # ~26 chunks, none mention the query terms
KEYWORD = b"Zygomorphic flowers have bilateral symmetry."


@pytest.fixture
def steered_vectors(monkeypatch: pytest.MonkeyPatch) -> None:
    """Vector search prefers every filler chunk; the keyword chunk is orthogonal to the question."""

    def embed(texts: list[str], *, model: str, timeout: float = 0.0) -> list[list[float]]:
        return [[0.0, 1.0, 0.0] if "Zygomorphic" in t else [1.0, 0.0, 0.0] for t in texts]

    monkeypatch.setattr("app.llm.embed", embed)


def _keyword_folder() -> int:
    return _built("Botany", ("filler.txt", FILLER), ("flowers.txt", KEYWORD))


def test_vector_only_misses_the_keyword_chunk(chat: FakeChat, steered_vectors: None) -> None:
    _keyword_folder()
    plan, _ = service.ask("zygomorphic flowers?")
    assert "flowers.txt" not in {c.filename for c in plan.chunks}
    assert {c.retriever for c in plan.chunks} == {"vector"}


def test_hybrid_surfaces_the_keyword_chunk_and_says_which_retriever(
    chat: FakeChat, steered_vectors: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("FILES_HYBRID", "1")
    _keyword_folder()
    plan, _ = service.ask("zygomorphic flowers?")
    assert len(plan.chunks) == 3
    keyword = next(c for c in plan.chunks if c.filename == "flowers.txt")
    assert keyword.retriever == "bm25"
    assert keyword.score == 0.0  # still the cosine similarity, not the fused score
    assert {c.retriever for c in plan.chunks if c.filename == "filler.txt"} == {"vector"}
    assert "Zygomorphic flowers" in chat.calls[-1]["prompt"]


def test_hybrid_marks_chunks_both_retrievers_found(chat: FakeChat, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FILES_HYBRID", "1")
    _built("Small", ("flowers.txt", KEYWORD), ("other.txt", b"Mosses reproduce with spores."))
    plan, _ = service.ask("zygomorphic flowers?")
    assert plan.chunks[0].filename == "flowers.txt" and plan.chunks[0].retriever == "both"


def test_hybrid_respects_folder_scope(chat: FakeChat, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FILES_HYBRID", "1")
    _built("Botany", ("flowers.txt", KEYWORD))
    other = _built("Other", ("notes.txt", b"Nothing about plants here."))
    plan, _ = service.ask("zygomorphic flowers?", folder_ids=[other])
    assert {c.folder_id for c in plan.chunks} == {other}


def test_bm25_index_follows_changes_in_the_folder(chat: FakeChat, monkeypatch: pytest.MonkeyPatch) -> None:
    from app.files import bm25

    folder_id = _built("Botany", ("other.txt", b"Mosses reproduce with spores."))
    assert bm25.search("zygomorphic", [folder_id], 5) == []
    service.add_files(folder_id, [("flowers.txt", KEYWORD)])
    service.build_folder(folder_id)
    assert len(bm25.search("zygomorphic", [folder_id], 5)) == 1


# --- 5. Conversational memory -------------------------------------------------------------


def test_memory_off_keeps_the_lab_prompt(chat: FakeChat) -> None:
    folder_id = _built("Lectures", ("a.txt", TXT))
    service.ask("first?", folder_ids=[folder_id])
    plan, _ = service.ask("second?", folder_ids=[folder_id])
    assert plan.memory is None
    assert chat.calls[-1]["prompt"] == answer.build_prompt("second?", plan.chunks)
    assert "Earlier in this conversation" not in chat.calls[-1]["prompt"]


def test_memory_includes_the_last_turns_from_the_same_scope(chat: FakeChat, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FILES_MEMORY_TURNS", "2")
    lectures = _built("Lectures", ("a.txt", TXT))
    labs = _built("Labs", ("b.txt", TXT + b" labs"))
    for q in ("one?", "two?", "three?"):
        service.ask(q, folder_ids=[lectures])
    service.ask("other scope?", folder_ids=[labs])

    plan, _ = service.ask("follow-up?", folder_ids=[lectures])
    assert plan.memory == [("two?", "Grounded answer."), ("three?", "Grounded answer.")]
    prompt = chat.calls[-1]["prompt"]
    assert "Q: two?\nA: Grounded answer.\nQ: three?\nA: Grounded answer." in prompt
    assert "one?" not in prompt and "other scope?" not in prompt
    # Memory sits apart from, and before, the retrieved context; the ONLY-context rule stays.
    assert prompt.index("Earlier in this conversation") < prompt.index("Context:")
    assert "Use ONLY the context below." in prompt and "must come from the Context" in prompt


# --- 6. Model picker ----------------------------------------------------------------------


def test_question_can_name_another_model(chat: FakeChat, client: TestClient) -> None:
    _built("Lectures", ("a.txt", TXT))
    client.post("/files/ask", json={"question": "q", "model": "qwen3:4b"})
    assert chat.calls[-1]["model"] == "qwen3:4b"
    assert client.get("/files/asks").json()[0]["model"] == "qwen3:4b"


def test_model_picker_off_always_uses_the_files_model(chat: FakeChat, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("FILES_MODEL_PICKER", "0")
    _built("Lectures", ("a.txt", TXT))
    service.ask("q", model="qwen3:4b")
    assert chat.calls[-1]["model"] == "phi3"
    assert service.chat_history()[0].model == "phi3"


def test_models_lists_pulled_chat_models_with_the_default_first(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.llm import OllamaModel

    pulled = [
        OllamaModel(name="nomic-embed-text:latest", family="nomic-bert"),
        OllamaModel(name="qwen3:4b", family="qwen3"),
        OllamaModel(name="phi3:latest", family="phi3"),
        OllamaModel(name="llama3:latest", family="llama"),
    ]
    monkeypatch.setattr("app.llm.list_models", lambda: pulled)
    assert client.get("/files/models").json() == {
        "default": "phi3",
        "models": ["phi3:latest", "llama3:latest", "qwen3:4b"],
        "error": None,
    }


def test_models_falls_back_to_the_default_when_ollama_is_down(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app import llm

    def down() -> list[Any]:
        raise llm.LlmError("connection refused")

    monkeypatch.setattr("app.llm.list_models", down)
    body = client.get("/files/models").json()
    assert body["models"] == ["phi3"] and body["error"] == "connection refused"


# --- settings -------------------------------------------------------------------------------


def test_settings_report_the_lab_defaults(client: TestClient) -> None:
    assert client.get("/files/settings").json() == {
        "accepted": [".pdf", ".txt", ".docx"],
        "default_model": "phi3",
        "top_k": 3,
        "show_sources": True,
        "hybrid": False,
        "memory_turns": 0,
        "model_picker": True,
        "score_threshold": None,
    }
